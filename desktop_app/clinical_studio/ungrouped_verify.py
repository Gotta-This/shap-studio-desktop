"""Native acceptance for unnamed vertical datasets, editing and group-free export."""
import json,traceback
from pathlib import Path
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication
from .merged_table_verify import save_fixture
from .typed_field_ui import FieldDialog
from .word_batch_engine import generate_batch,export_batch


def verify_ungrouped(window,directory):
    root=Path(directory).resolve();root.mkdir(parents=True,exist_ok=True);app=QApplication.instance();checks=[]
    def check(name,value):
        checks.append(dict(name=name,passed=bool(value)));print(name,bool(value),flush=True)
        if not value:raise AssertionError(name)
    def pump():
        for _ in range(4):app.processEvents()
    report=dict(passed=False)
    try:
        window.timer.stop();tool=window.open_module('word')
        rows=[['项目','总体(n=20)'],['性别（男/女）','18(90)/2(10)'],['年龄（岁）','58.35±10.41'],['BMI（kg/m²）','23.85±2.74'],['合并高血压','7(35)']]
        path=save_fixture(root/'ungrouped.docx',[rows]);tool.import_word(path);pump()
        check('无组别文档直接识别四个字段',tool.word_scan['candidates']==[] and len(tool.preview_draft['fields'])==4)
        check('无分组模式不要求填写组名',tool.ungrouped_option.isChecked() and not tool.group_area.isVisible() and tool.group_inputs[0].text()=='')
        check('原表例数自动带入',tool.ungrouped_n.value()==20)
        window.grab().save(str(root/'ungrouped-import.png'))
        errors=[]
        def edit_preview():
            dialog=tool.scan_dialog
            try:
                check('预览将归属列显示为数据集',dialog.fields.typed_table.horizontalHeaderItem(2).text()=='数据集')
                index=next(i for i,f in enumerate(dialog.fields.typed_specs) if f['name']=='年龄（岁）')
                def edit_value():
                    editor=dialog.fields.field_dialog
                    try:
                        check('无分组字段编辑为原生窗口',isinstance(editor,FieldDialog) and editor.graphicsProxyWidget() is None)
                        check('单数据集不显示跨组同步',not editor.sync.isVisible())
                        editor.inputs['mean'].setText('61.23');editor.apply()
                    except Exception:errors.append(traceback.format_exc());editor.reject()
                QTimer.singleShot(100,edit_value);dialog.fields.edit_typed_field(index)
                dialog.tabs.setCurrentIndex(1);pump();dialog.grab().save(str(root/'ungrouped-preview.png'));dialog.accept()
            except Exception:errors.append(traceback.format_exc());dialog.reject()
        QTimer.singleShot(120,edit_preview);tool.open_scan_preview();pump()
        if errors:raise AssertionError(errors[0])
        tool.confirm_groups_button.click();pump();check('不用组名即可应用识别',tool.ungrouped and tool.groups_confirmed)
        check('预览修改保留',next(f for f in tool.typed_specs if f['name']=='年龄（岁）')['mean']==61.23)
        config=tool.get_config();result=generate_batch(config)
        check('生成20例且不含分组列',len(result['rows'])==20 and '分组' not in result['columns'])
        check('男女人数按原表保留',sum(row[result['columns'].index('性别')]=='男' for row in result['rows'])==18)
        output=root/'ungrouped-cases.xlsx';export_batch(result,output)
        from openpyxl import load_workbook
        book=load_workbook(output);check('Excel没有虚构分组列且有20条记录',book.worksheets[0].max_row==21 and '分组' not in [cell.value for cell in book.worksheets[0][1]]);book.close()
        tool.common_n.setValue(30);pump();larger=generate_batch(tool.get_config())
        check('生成例数可变且保留原比例',len(larger['rows'])==30 and sum(row[larger['columns'].index('性别')]=='男' for row in larger['rows'])==27)
        window.grab().save(str(root/'ungrouped-applied.png'))
        missing=save_fixture(root/'missing-n.docx',[[['指标','数值'],['年龄','60±5'],['BMI','23±2']]])
        tool.import_word(missing);pump();check('未报告例数仍先识别字段',len(tool.preview_draft['fields'])==2 and tool.ungrouped_n.value()==0)
        check('不使用默认例数直接生成',not tool.generate_button.isEnabled() and not tool.confirm_groups_button.isEnabled())
        window.grab().save(str(root/'missing-n.png'))
        tool.ungrouped_n.setValue(12);tool.confirm_groups();pump()
        check('补充例数后无组别也能生成',tool.ungrouped and len(generate_batch(tool.get_config())['rows'])==12)
        report['passed']=True
    except Exception:report['traceback']=traceback.format_exc()
    report['checks']=checks;(root/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');return report['passed']
