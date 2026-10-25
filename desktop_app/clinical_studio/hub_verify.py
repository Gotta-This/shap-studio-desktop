"""Operate all four packaged modules and verify their independent exports."""
from collections import Counter
from datetime import datetime
import csv
import json
from pathlib import Path
import time
import traceback
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from PySide6.QtCore import Qt, QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QLineEdit

from .ui_verify import UiVerifier


def make_word_fixture(path):
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    root = ET.Element(ns+'document')
    body = ET.SubElement(root, ns+'body')
    table = ET.SubElement(body, ns+'tbl')
    for row in [('组别','例数','测试指标'),('对照组','10','10.00±2.00'),('中药组','12','12.00±3.00')]:
        tr = ET.SubElement(table, ns+'tr')
        for value in row:
            tc = ET.SubElement(tr, ns+'tc')
            p = ET.SubElement(tc, ns+'p')
            r = ET.SubElement(p, ns+'r')
            ET.SubElement(r, ns+'t').text = value
    with ZipFile(path,'w') as archive:
        archive.writestr('word/document.xml',ET.tostring(root,encoding='utf-8',xml_declaration=True))
        archive.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr('_rels/.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')


class SuiteVerifier:
    ensure = UiVerifier.ensure
    click = UiVerifier.click
    type_text = UiVerifier.type_text
    combo = UiVerifier.combo
    tab = UiVerifier.tab
    save_dialog = UiVerifier.save_dialog

    def __init__(self,hub,directory,app):
        self.hub,self.app=hub,app
        self.root=Path(directory).resolve();self.root.mkdir(parents=True,exist_ok=True)
        self.shots=self.root/'screenshots';self.shots.mkdir(exist_ok=True)
        self.events=[];self.checks=[];self.failed=False;self.completed=False;self.shap=None
        self.modifier=Qt.ControlModifier

    @property
    def passed(self):return self.completed and not self.failed and self.shap is not None and self.shap.passed

    def log(self,message,**details):self.events.append(dict(action=message,**details))

    def check(self,label,condition,**details):
        self.checks.append(dict(item=label,passed=bool(condition),**details))
        if not condition:raise AssertionError(label+': '+str(details))

    def screenshot(self,name):
        self.app.processEvents();self.hub.grab().save(str(self.shots/(name+'.png')))

    def step(self,fn):
        try:fn()
        except Exception as exc:self.fail(exc)

    def fail(self,exc):
        if self.failed:return
        self.failed=True;self.log('四模块验证失败',error=str(exc),traceback=traceback.format_exc())
        self.report(False)
        (self.root/'ui_verification.json').write_text(json.dumps(dict(passed=False,events=self.events,checks=self.checks),ensure_ascii=False,indent=2),encoding='utf-8')
        QTimer.singleShot(100,self.app.quit)

    def report(self,passed):
        (self.root/'four_module_verification.json').write_text(json.dumps(dict(passed=passed,time=datetime.now().isoformat(),checks=self.checks,events=self.events),ensure_ascii=False,indent=2),encoding='utf-8')

    def enter(self,key):
        if self.hub.current_module is not None:self.click(self.hub.home_button)
        self.click(self.hub.launch_buttons[key])
        self.check('独立入口：'+key,self.hub.current_module==key)
        return self.hub.modules[key]

    def table_value(self,table,row,col,value):
        cell=table.item(row,col)
        QTest.mouseClick(table.viewport(),Qt.LeftButton,pos=table.visualItemRect(cell).center())
        table.setFocus();table.setCurrentItem(cell);table.editItem(cell)
        self.app.processEvents()
        editor=next((e for e in table.findChildren(QLineEdit) if e.isVisible()),None)
        if not isinstance(editor,QLineEdit):raise AssertionError('未打开类别表格编辑框')
        self.type_text(editor,value);QTest.keyClick(table,Qt.Key_Return);self.app.processEvents()

    def start(self):self.step(self.categories)

    def categories(self):
        self.check('首页恰有四个独立入口',len(self.hub.launch_buttons)==4 and not self.hub.modules)
        self.screenshot('工具集_00_四模块首页')
        tool=self.enter('category')
        self.check('只创建已进入的类别工具',set(self.hub.modules)=={'category'})
        self.table_value(tool.table,0,1,'11');self.table_value(tool.table,1,1,'9')
        self.type_text(tool.seed,'44');self.click(tool.generate_button)
        self.check('类别数列各类数量与界面一致',Counter(tool.result.values)=={1:11,0:9})
        self.click(tool.result.copy_button)
        self.check('类别数列复制与显示一致',QApplication.clipboard().text()==tool.result.output.toPlainText())
        path=self.root/'工具1_类别数列.csv';self.save_dialog(tool.result.save_button,path)
        with path.open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream))
        self.check('类别 CSV 全部数值与界面一致',rows[0]==['RandomValue'] and [float(r[0]) for r in rows[1:]]==tool.result.values)
        self.category_snapshot=(tool.seed.text(),tuple(tool.result.values))
        self.screenshot('工具集_01_类别数列独立页面')
        QTimer.singleShot(100,lambda:self.step(self.numeric))

    def numeric(self):
        tool=self.enter('numeric')
        for field,value in [(tool.mean,'10'),(tool.sd,'2'),(tool.lower,'0'),(tool.upper,'20'),(tool.seed,'42')]:self.type_text(field,value)
        self.type_text(tool.n.findChild(QLineEdit),'32');QTest.keyClick(tool.n,Qt.Key_Return)
        self.combo(tool.digits,'2');self.click(tool.generate_button)
        self.wait_for(tool,self.numeric_done)

    def wait_for(self,tool,callback,deadline=None):
        if self.failed:return
        deadline=deadline or time.monotonic()+60
        if tool.is_running():
            if time.monotonic()>deadline:self.fail(TimeoutError('工具后台任务超时'));return
            QTimer.singleShot(100,lambda:self.wait_for(tool,callback,deadline));return
        self.app.processEvents();self.step(callback)

    def numeric_done(self):
        tool=self.hub.modules['numeric']
        self.check('数值工具例数、范围与界面一致',len(tool.result.values)==32 and all(0<=v<=20 for v in tool.result.values),message=tool.error.text())
        path=self.root/'工具2_均数标准差.xlsx';self.save_dialog(tool.result.save_button,path)
        from openpyxl import load_workbook
        book=load_workbook(path,data_only=True)
        values=[r[0] for r in list(book.active.values)[1:]];book.close()
        self.check('数值 Excel 全部数值与界面一致',values==tool.result.values)
        self.numeric_snapshot=(tool.mean.text(),tool.sd.text(),tool.n.value(),tool.seed.text(),tuple(tool.result.values))
        self.screenshot('工具集_02_均数标准差独立页面')
        QTimer.singleShot(100,lambda:self.step(self.word))

    def word(self):
        tool=self.enter('word')
        source=self.root/'工具3_测试Word.docx';make_word_fixture(source)
        tool.use_word_sizes.setChecked(True)
        self.save_dialog(tool.import_button,source)
        self.check('Word 经导入对话框识别两组指标',tool.metrics.rowCount()==2 and tool.group_editors[0].fields['n'].text()=='10',message=tool.message.text())
        self.tab(tool.pages,0);tool.same_size.setChecked(False);self.type_text(tool.seed,'77')
        for index,count in enumerate(('14','16')):
            tool.group_selector.setCurrentIndex(index);self.app.processEvents();self.type_text(tool.group_editors[index].fields['n'],count)
        self.click(tool.generate_button);self.wait_for(tool,self.word_done)

    def word_done(self):
        tool=self.hub.modules['word']
        self.check('Word 工具按修改后的两组例数生成',tool.result is not None and len(tool.result['rows'])==30,message=tool.message.text())
        self.check('Word 两组人数与配置一致',Counter(r[1] for r in tool.result['rows'])=={'对照组':14,'中药组':16})
        self.word_export=self.root/'工具3_Word生成Excel.xlsx'
        self.save_dialog(tool.export_button,self.word_export,must_exist=False)
        self.wait_for(tool,self.word_saved)

    def word_saved(self):
        tool=self.hub.modules['word']
        from openpyxl import load_workbook
        book=load_workbook(self.word_export,data_only=True)
        rows=list(book['模拟病例数据'].values)
        same=rows[0]==tuple(tool.result['columns']) and len(rows)==len(tool.result['rows'])+1
        for exported,expected in zip(rows[1:],tool.result['rows']):
            same=bool(same and all(abs(a-b)<1e-10 if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==(None if b=='' else b) for a,b in zip(exported,expected)))
        self.check('Word 导出 Excel 全部单元与界面结果一致',same,sheets=book.sheetnames)
        book.close()
        self.word_snapshot=json.dumps(dict(config=tool.get_config(),result=tool.result),ensure_ascii=False,sort_keys=True)
        self.screenshot('工具集_03_Word独立页面')
        QTimer.singleShot(100,lambda:self.step(self.start_shap))

    def start_shap(self):
        tool=self.enter('shap')
        self.check('生图工具仅有自己的六个导航项',tool.nav.count()==6 and not hasattr(tool,'data_tools_page'))
        self.check('四个模块分别持有不同实例',len({id(w) for w in self.hub.modules.values()})==4)
        self.screenshot('工具集_04_模拟数据生图独立页面')
        self.shap=UiVerifier(tool,self.root,self.app)
        original=self.shap.finish_report
        def finish(passed):
            original(passed)
            if not passed:self.report(False);return
            try:
                c=self.hub.modules['category'];n=self.hub.modules['numeric'];w=self.hub.modules['word']
                self.check('生图操作未改变类别工具参数与结果',(c.seed.text(),tuple(c.result.values))==self.category_snapshot)
                self.check('生图操作未改变数值工具参数与结果',(n.mean.text(),n.sd.text(),n.n.value(),n.seed.text(),tuple(n.result.values))==self.numeric_snapshot)
                self.check('生图操作未改变 Word 参数与结果',json.dumps(dict(config=w.get_config(),result=w.result),ensure_ascii=False,sort_keys=True)==self.word_snapshot)
                self.completed=True;self.report(True)
            except Exception as exc:self.fail(exc)
        self.shap.finish_report=finish
        self.shap.start()
