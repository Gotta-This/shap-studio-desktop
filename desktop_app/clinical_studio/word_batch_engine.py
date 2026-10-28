"""Import Word summary tables and create explicitly labelled synthetic case data.

DOCX parsing uses the OOXML grid (including horizontal/vertical merges), not
positional guesses based on the number of physical cells. No source files change.
"""
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass, field
import json
import math
import os
from pathlib import Path
import random
import re
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile

from .data_tools_engine import numeric_numbers, statistics

GROUPS = ('对照组', '中药组')
BASE_COLUMNS = ['编号', '分组', '性别', '年龄', '病程', '病情程度', '病变范围']
END_COLUMNS = ['总有效率', '复发率']
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
NUM = r'[+−-]?(?:\d+(?:\.\d*)?|\.\d+)'
PAIR = re.compile(r'('+NUM+r')\s*±\s*('+NUM+r')')


@dataclass
class ImportResult:
    indicators: list
    group_sizes: dict
    outcome_rates: dict
    messages: list
    source: str
    table_count: int
    group_matches: dict = field(default_factory=dict)
    table_matches: list = field(default_factory=list)
    general_groups: dict = field(default_factory=dict)
    text_group_matches: dict = field(default_factory=dict)
    typed_fields: list = field(default_factory=list)
    table_issues: list = field(default_factory=list)
    baseline_columns: bool = False


def table_titles(body,tables):
    titles={};recent=''
    for child in body:
        if child.tag==W+'p':
            text=''.join(t.text or '' for t in child.iter(W+'t')).strip()
            if re.match(r'^表\s*[\d一二三四五六七八九十]+',text):recent=text
            elif text:recent=''
        elif child.tag==W+'tbl':
            for table in child.iter(W+'tbl'):
                if table in tables:titles[tables.index(table)+1]=recent or f'表{tables.index(table)+1}（未提供表名）'
            recent=''
    return titles


def _grid(table, with_origins=False):
    rows = []
    origins = []
    above = []
    for row_index,tr in enumerate(table.findall(W+'tr')):
        row = []
        origin = []
        before = tr.find(W+'trPr/'+W+'gridBefore')
        if before is not None:
            row.extend(['']*int(before.get(W+'val', '0')))
            origin.extend([None]*len(row))
        for cell_index,cell in enumerate(tr.findall(W+'tc')):
            text = ''.join(t.text or '' for t in cell.iter(W+'t')).strip()
            span = cell.find(W+'tcPr/'+W+'gridSpan')
            width = int(span.get(W+'val', '1')) if span is not None else 1
            merge = cell.find(W+'tcPr/'+W+'vMerge')
            if merge is not None and merge.get(W+'val') != 'restart' and not text:
                text = above[len(row)] if len(row) < len(above) else ''
            row.extend([text]*width)
            origin.extend([(row_index,cell_index)]*width)
        rows.append(row)
        origins.append(origin)
        above = row
    return (rows,origins) if with_origins else rows


def _clean(value):
    return re.sub(r'\s+', '', value).replace('−', '-')


def _header(headers, col):
    labels = []
    for row in headers:
        value = row[col].strip() if col < len(row) else ''
        if value and value not in labels:
            labels.append(value)
    phase = next(('前' if '前' in x else '后' for x in reversed(labels)
                  if x in ('治疗前', '治疗后', '干预前', '干预后', '前', '后')), '')
    metric = next((x for x in labels if x not in ('治疗前', '治疗后', '干预前', '干预后', '前', '后')), '')
    metric = re.sub(r'[（(].*?[）)]', '', metric).strip()
    return (phase+'-' if phase else '')+metric


def group_names(names):
    names=[str(name).strip() for name in names]
    if not 1<=len(names)<=20 or any(not n or len(n)>100 for n in names):
        raise ValueError('请填写1～20个组别，每个名称须为1～100个字符。')
    if len({_clean(n) for n in names})!=len(names):
        raise ValueError('组别名称不能重复（忽略空格和换行）。')
    return names


def read_word(path, groups=None, *, parse_indicators=True):
    groups=list(GROUPS) if groups is None else list(groups)
    if groups:groups=group_names(groups)
    elif parse_indicators:raise ValueError('请先填写并确认组别。')
    original_path=Path(path).resolve()
    from .word_formats import ensure_docx
    path=ensure_docx(original_path)
    try:
        with ZipFile(path) as archive:
            info = archive.getinfo('word/document.xml')
            if info.file_size > 25_000_000:
                raise ValueError('Word 正文过大（上限 25 MB），请精简文件后导入。')
            root = ET.fromstring(archive.read(info))
    except (BadZipFile, KeyError, ET.ParseError) as exc:
        raise ValueError('无法读取 Word 正文，请确认文件为有效的 .docx。') from exc
    body = root.find(W+'body')
    if body is None:
        raise ValueError('Word 中没有正文。')
    tables = list(body.iter(W+'tbl'))
    titles=table_titles(body,tables)
    from .word_prose import read_general_prose
    general_groups,text_group_matches=read_general_prose(body,groups)
    indicators, sizes, rates, messages = [], {}, {}, []
    used_names = set(BASE_COLUMNS+END_COLUMNS)
    group_matches={g:0 for g in groups};table_matches=[]
    typed_fields=[];table_issues=[];baseline_columns=False
    from .word_table_types import parse_table
    for index, table in enumerate(tables, 1):
        rows,origins = _grid(table,True)
        typed=parse_table(rows,origins,groups,index,table_title=titles.get(index,''))
        group_rows = {g: [row for row in rows if any(_clean(cell)==_clean(g) for cell in row)] for g in groups}
        matched={g:len(rs) for g,rs in group_rows.items() if rs}
        matched.update(typed['matches'])
        table_matches.append(dict(table=index,groups=matched))
        for g,count in matched.items():group_matches[g]+=count
        if not parse_indicators:continue
        typed_fields.extend(typed['fields']);table_issues.extend(typed['issues']);sizes.update(typed['sizes']);baseline_columns=baseline_columns or typed['baseline']
        for group,values in typed.get('rates',{}).items():rates.setdefault(group,{}).update(values)
        if typed['handled']:
            messages.append(f'表 {index}：识别 {len(typed["fields"])} 个组别字段（支持竖表及多层表头）；字段类型和分类可修改。')
            continue
        if not any(group_rows.values()):
            messages.append(f'表 {index}：未找到所选组别数据行，未导入。')
            continue
        if any(len(v)>1 for v in group_rows.values()):
            messages.append(f'表 {index}：同组存在多行，无法自动确定，整表未导入；请人工补录。')
            continue
        first = min(i for i,row in enumerate(rows) if any(row in rs for rs in group_rows.values()))
        headers = rows[:first]
        for group, candidates in group_rows.items():
            if not candidates:
                messages.append(f'表 {index}：缺少{group}行；相关指标会标记为待补全。')
                continue
            row = candidates[0]
            for col,value in enumerate(row):
                count_header = _header(headers,col).lower()
                if ('例数' in count_header or count_header in ('n','样本量')) and _clean(value).isdigit():
                    n = int(_clean(value))
                    if group in sizes and sizes[group]!=n:
                        messages.append(f'表 {index}：{group}例数 {n} 与此前 {sizes[group]} 不同，请核对配置。')
                    else:sizes[group] = n
        width = max(map(len, rows), default=0)
        found = 0
        relevant=[i for i,row in enumerate(rows) if i<first or any(row in rs for rs in group_rows.values())]
        for col in range(width):
            # Word sometimes splits one physical data cell across several grid columns.
            # Import it once when the headers AND both data rows share that same span.
            if col>1 and all(col<len(origins[i]) and origins[i][col] is not None and origins[i][col]==origins[i][col-1] for i in relevant):
                continue
            raw = {g: (rs[0][col] if rs and col<len(rs[0]) else '') for g, rs in group_rows.items()}
            matches = {g: PAIR.search(_clean(v)) for g, v in raw.items()}
            name = _header(headers, col)
            if any(matches.values()):
                if not name:
                    name = f'表{index}-指标{col-1}'
                    messages.append(f'表 {index}第 {col+1} 列：无可识别表头，已使用临时名称。')
                if name in used_names:
                    name += f'（表{index}列{col+1}）'
                used_names.add(name)
                for group in groups:
                    match = matches[group]
                    if not match:
                        messages.append(f'表 {index} {name} {group}：未识别均数±标准差，生成前请补全或移除该指标的全部组别行。')
                    indicators.append(dict(table=index, name=name, group=group,
                                           mean=float(match[1]) if match else None,
                                           sd=float(match[2]) if match else None,
                                           digits=2, lower=None, upper=None, parity='any', raw=raw[group]))
                found += 1
            elif name in END_COLUMNS:
                for group, value in raw.items():
                    value = _clean(value).replace('（','(').replace('）',')')
                    match = re.match(r'(\d+)\((\d+(?:\.\d+)?)%?\)', value)
                    if match and sizes.get(group):
                        rates.setdefault(group, {})[name] = int(match[1])/sizes[group]*100
                    else:
                        pct = re.fullmatch(r'(\d+(?:\.\d+)?)%', value)
                        if pct:
                            rates.setdefault(group, {})[name] = float(pct[1])
                        else:
                            messages.append(f'表 {index} {group} {name}：比例未识别，请手动设置。')
        messages.append(f'表 {index}：识别 {found} 个数值指标。')
    if parse_indicators and not indicators and not typed_fields and not any(data['indicators'] for data in general_groups.values()):
        messages.append('未识别到均数±标准差指标，可在指标配置中手动添加。')
    for group,data in general_groups.items():
        if group in sizes and sizes[group]!=data['n']:
            data['issues'].append('正文例数与表格例数不一致')
        else:sizes[group]=data['n']
        messages.append(f'正文 {group}：'+data['raw'])
        for issue in data['issues']:messages.append(f'正文 {group}：{issue}，请修正后重新导入。')
    messages.append('正文一般资料按固定句式识别；未识别的基本资料仍需人工配置，不能视为Word中的参数。')
    messages.append('各指标按列独立模拟，不还原病例配对、相关性或原文检验结果。')
    messages.extend(table_issues)
    return ImportResult(indicators, sizes, rates, messages, str(original_path), len(tables), group_matches, table_matches,general_groups,text_group_matches,typed_fields,table_issues,baseline_columns)


def default_batch_config(groups=GROUPS):
    return dict(version=1, seed=12, source='', course_unit='月', groups=[dict(
        name=g, n=70, male=50., severity=[50.,50.,0.], effective=80., recurrence=10.,
        age=[dict(lower=18, upper=74, percent=100.)],
        course=[dict(lower=1, upper=60, percent=100.)], course_digits=0,
        extent_mode='blank', recurrence_scope='all',
        extent=[dict(label='自定义类别',percent=100.)],
        extent_ranges=[dict(lower=1,upper=3,percent=100.)],extent_digits=0)
        for g in group_names(groups)], indicators=[])


def allocation(n, weights):
    if not weights or any(not math.isfinite(w) or w<0 for w in weights) or abs(math.fsum(weights)-100)>1e-6:
        raise ValueError('每组比例须为非负数，且合计为 100%。')
    expected = [n*w/100 for w in weights]
    counts = [math.floor(v) for v in expected]
    for i in sorted(range(len(weights)),key=lambda i: (-(expected[i]-counts[i]),i))[:n-sum(counts)]:
        counts[i] += 1
    return counts


def _finite(value, label):
    if isinstance(value, bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
        raise ValueError(f'{label}须为有效数字。')
    return value


def validate_batch(config):
    if config.get('version') != 1:
        raise ValueError('不支持此配置版本。')
    seed = config.get('seed')
    if not isinstance(seed,int) or isinstance(seed,bool) or not 0<=seed<=4294967295:
        raise ValueError('随机种子须为 0～4294967295 的整数。')
    groups = config.get('groups',[])
    selected_groups=group_names([g.get('name','') for g in groups])
    if config.get('ungrouped') and (len(groups)!=1 or config.get('include_basics',True)):
        raise ValueError('无分组模式仅支持一个数据集，且只输出已识别字段。')
    if any(g['name']!=name for g,name in zip(groups,selected_groups)):
        raise ValueError('分组名称首尾不能含空格。')
    for g in groups:
        if not isinstance(g.get('blank_fields',[]),list) or any(name not in BASE_COLUMNS[2:]+END_COLUMNS for name in g.get('blank_fields',[])):
            raise ValueError('留空字段配置无效。')
        n=g.get('n')
        if not isinstance(n,int) or isinstance(n,bool) or not 2<=n<=5000:
            raise ValueError(f"{g['name']}例数须为 2～5000 的整数。")
        for name in ('male','effective','recurrence'):
            if not 0<=_finite(g[name],name)<=100:
                raise ValueError(f'{name}须在 0～100% 之间。')
        if len(g['severity'])!=3:
            raise ValueError('病情程度须包含轻度、中度、重度三个比例。')
        allocation(n,[_finite(x,'病情比例') for x in g['severity']])
        for key in ('age','course'):
            ranges=g[key]
            allocation(n,[_finite(r['percent'],'区间比例') for r in ranges])
            d=0 if key=='age' else g['course_digits']
            if not isinstance(d,int) or d not in range(5):
                raise ValueError('病程小数位须为 0～4。')
            for r in ranges:
                low,high=_finite(r['lower'],'区间下限'),_finite(r['upper'],'区间上限')
                if low<0 or low>high or high>1e6 or math.ceil(low*10**d)>math.floor(high*10**d):
                    raise ValueError('年龄／病程区间须非负、有可用数值，且上限不超过 100 万。')
            for r,lo,hi in _effective_ranges(ranges,d):
                if lo>hi:raise ValueError('区间在当前小数位下没有可用数值。')
            active=sorted(ranges,key=lambda r:r['lower'])
            if any(a['upper']>b['lower'] for a,b in zip(active,active[1:])):
                raise ValueError('年龄／病程区间不能交叉；相邻区间允许共用边界，例如10～20、20～30，边界值归入后一段。')
        if g.get('age_stats') is not None:
            if len(g['age'])!=1:raise ValueError('按年龄均数/标准差生成时，只能设置一个年龄范围。')
            mean=_finite(g['age_stats']['mean'],'年龄均数');sd=_finite(g['age_stats']['sd'],'年龄标准差')
            lo,hi=g['age'][0]['lower'],g['age'][0]['upper']
            if not lo<=mean<=hi or sd<0:raise ValueError(f"{g['name']}年龄均数须在范围内，标准差须非负。")
            maximum=math.sqrt(max(0,(mean-lo)*(hi-mean)*n/(n-1)))
            if sd>maximum+1e-9:raise ValueError(f"{g['name']}年龄范围 {lo:g}～{hi:g}、均数 {mean:g} 与标准差 {sd:g} 无法同时满足；该范围和均数下样本标准差最多为 {maximum:.2f}，请核对Word原文。")
        if g.get('recurrence_scope','all') not in ('all','effective'):
            raise ValueError('复发率分母须为全组或治疗有效者。')
        extent_mode=g.get('extent_mode','category')
        if extent_mode=='category':
            extent=g['extent'];allocation(n,[_finite(r['percent'],'病变范围比例') for r in extent])
            labels=[r['label'].strip() for r in extent]
            if any(not x for x in labels) or len(labels)!=len(set(labels)):
                raise ValueError('病变范围类别不能为空或重复。')
        elif extent_mode=='range':
            ranges=g.get('extent_ranges',[]);allocation(n,[_finite(r['percent'],'病变范围区间比例') for r in ranges])
            d=g.get('extent_digits',0)
            if not isinstance(d,int) or d not in range(5):raise ValueError('病变范围小数位须为0～4。')
            for r in ranges:
                lo,hi=_finite(r['lower'],'病变范围下限'),_finite(r['upper'],'病变范围上限')
                if max(abs(lo),abs(hi))>1e6 or lo>hi or math.ceil(lo*10**d)>math.floor(hi*10**d):raise ValueError('病变范围区间无有效数值。')
            if any(lo>hi for r,lo,hi in _effective_ranges(ranges,d)):raise ValueError('病变范围区间没有可用数值。')
            ordered=sorted(ranges,key=lambda r:r['lower'])
            if any(x['upper']>y['lower'] for x,y in zip(ordered,ordered[1:])):raise ValueError('病变范围区间不能重叠。')
        elif extent_mode!='blank':raise ValueError('病变范围模式无效。')
    categorical_specs=config.get('categoricals',[]);category_names={}
    for spec in categorical_specs:
        name=spec['name'].strip();group=spec['group']
        if not name or name in BASE_COLUMNS+END_COLUMNS or name.startswith('='):raise ValueError('正文分类字段名称无效或与固定列重名。')
        if group not in selected_groups or group in category_names.setdefault(name,set()):raise ValueError('正文分类字段组别无效或重复。')
        category_names[name].add(group)
        labels=spec['labels']
        if not labels or any(not str(v).strip() for v in labels) or len(set(labels))!=len(labels) or len(labels)!=len(spec['weights']):raise ValueError('正文分类字段类别为空、重复或比例不完整。')
        allocation(groups[0]['n'],[_finite(v,'正文类别比例') for v in spec['weights']])
    if any(gs!=set(selected_groups) for gs in category_names.values()):raise ValueError('每个正文分类字段须配置全部所选组别。')
    specs=config.get('indicators',[])
    if not specs and not categorical_specs and not config.get('typed_fields') and not any(g.get('age_stats') for g in groups) and not config.get('prose_sources'):
        raise ValueError('请导入 Word 或添加至少一个数值指标。')
    if len(specs)>250*len(selected_groups):
        raise ValueError('最多支持250个指标，每个指标须配置所有所选组别。')
    names={}
    for spec in specs:
        name=spec['name'].strip()
        if not name or name in BASE_COLUMNS+END_COLUMNS or name in category_names or name.startswith('='):
            raise ValueError('指标名不能为空、与固定列重名或以等号开头。')
        group=spec['group']
        if group not in selected_groups or group in names.setdefault(name,set()):
            raise ValueError(f'{name} 的分组无效或重复。')
        names[name].add(group)
        _finite(spec['mean'],name+'均数')
        if _finite(spec['sd'],name+'标准差')<0:
            raise ValueError(f'{name} 的标准差须非负。')
        digits=spec['digits']
        if not isinstance(digits,int) or digits not in range(5) or spec['parity'] not in ('any','odd','even'):
            raise ValueError('指标小数位须为 0～4，数值类型须有效。')
        if spec['parity']!='any' and digits!=0:
            raise ValueError('奇偶数指标的小数位须为 0。')
        for key in ('lower','upper'):
            if spec.get(key) is not None:
                _finite(spec[key],name+'边界')
        if spec.get('lower') is not None and spec.get('upper') is not None and spec['lower']>spec['upper']:
            raise ValueError(f'{name} 下限不能超过上限。')
    if any(gs!=set(selected_groups) for gs in names.values()):
        raise ValueError('每个指标都须配置全部所选组别；请补全缺失行。')
    if len(names)>250:raise ValueError('最多支持250个指标。')
    from .typed_field_engine import validate_fields
    validate_fields(config)


def output_columns(config):
    names=[s['name'].strip() for s in config.get('categoricals',[])+config.get('indicators',[])+config.get('typed_fields',[]) if s.get('kind')!='excluded']
    base=['编号'] if config.get('ungrouped') else BASE_COLUMNS if config.get('include_basics',True) else ['编号','分组']
    return list(dict.fromkeys(base+names+(END_COLUMNS if config.get('include_basics',True) else [])))


def _categorical(n, labels, weights, rng):
    a=[label for label,count in zip(labels,allocation(n,weights)) for _ in range(count)]
    rng.shuffle(a)
    return a


def _effective_ranges(ranges,digits):
    factor=10**digits
    ordered=sorted(enumerate(ranges),key=lambda x:x[1]['lower'])
    bounds={}
    for j,(i,r) in enumerate(ordered):
        low,high=math.ceil(r['lower']*factor),math.floor(r['upper']*factor)
        if j+1<len(ordered) and r['upper']==ordered[j+1][1]['lower']:
            high=math.ceil(r['upper']*factor)-1
        bounds[i]=(low,high)
    return [(r,*bounds[i]) for i,r in enumerate(ranges)]


def _ranges(n, ranges, digits, rng):
    factor=10**digits;a=[]
    for (r,lo,hi),count in zip(_effective_ranges(ranges,digits),allocation(n,[r['percent'] for r in ranges])):
        a.extend(rng.randint(lo,hi)/factor for _ in range(count))
    rng.shuffle(a);return a


def _refine_indicator(values,spec):
    """Reduce moment error after decimal rounding without violating bounds."""
    if spec['digits']==0:return values
    mean,sd=statistics(values)
    if spec['sd']==0:candidate=[round(spec['mean'],spec['digits'])]*len(values)
    elif sd>0:candidate=[round((v-mean)*spec['sd']/sd+spec['mean'],spec['digits']) for v in values]
    else:return values
    if any((spec.get('lower') is not None and v<spec['lower']) or
           (spec.get('upper') is not None and v>spec['upper']) for v in candidate):return values
    cm,cs=statistics(candidate)
    return candidate if (cm-spec['mean'])**2+(cs-spec['sd'])**2 < (mean-spec['mean'])**2+(sd-spec['sd'])**2 else values


def _indicator(spec,n,rng,cancelled=None):
    options=dict(n=n,mean=spec['mean'],sd=spec['sd'],digits=spec['digits'],
                 lower=spec.get('lower'),upper=spec.get('upper'),seed=rng.getrandbits(32))
    if spec['parity']=='any':
        return _refine_indicator(numeric_numbers(**options,cancelled=cancelled),spec)
    parity=1 if spec['parity']=='odd' else 0
    options.update(mean=(spec['mean']-parity)/2,sd=spec['sd']/2,digits=0,
                   lower=math.ceil((spec['lower']-parity)/2) if spec.get('lower') is not None else None,
                   upper=math.floor((spec['upper']-parity)/2) if spec.get('upper') is not None else None)
    return [2*x+parity for x in numeric_numbers(**options,cancelled=cancelled)]


def _age_values(spec,n,rng,cancelled):
    """Refine rounded integer ages while retaining the source range."""
    values=_indicator(spec,n,rng,cancelled);buckets={}
    for i,value in enumerate(values):buckets.setdefault(value,[]).append(i)
    total=sum(values);squares=sum(v*v for v in values)
    def error(s,q):
        mean=s/n;sd=math.sqrt(max(0,(q-s*s/n)/(n-1)))
        return (mean-spec['mean'])**2+(sd-spec['sd'])**2
    for iteration in range(min(5000,n*50)):
        if cancelled():raise InterruptedError('已取消生成。')
        score=error(total,squares);best=None
        for value in sorted(buckets):
            for delta in (-1,1):
                candidate=value+delta
                if not spec['lower']<=candidate<=spec['upper']:continue
                q=squares+2*value*delta+1;new_score=error(total+delta,q)
                if new_score<score-1e-12:score=new_score;best=(value,candidate,delta,q)
        if best is None:break
        value,candidate,delta,squares=best;index=buckets[value].pop()
        if not buckets[value]:del buckets[value]
        buckets.setdefault(candidate,[]).append(index);values[index]=candidate;total+=delta
    return values


def generate_batch(config, progress=lambda _:None, cancelled=lambda:False):
    config=deepcopy(config)
    validate_batch(config)
    rng=random.Random(config['seed'])
    names=list(dict.fromkeys(s['name'].strip() for s in config['indicators']))
    categorical_names=list(dict.fromkeys(s['name'].strip() for s in config.get('categoricals',[])))
    columns=output_columns(config)
    rows,checks=[],[]
    for g in config['groups']:
        n=g['n'];group=g['name']
        def categories(field,labels,weights):
            values=_categorical(n,labels,weights,rng)
            counts=Counter(values)
            for label,pct in zip(labels,weights):
                checks.append(dict(group=group,field=field,category=label,target_percent=pct,actual_count=counts[label],actual_percent=counts[label]/n*100))
            return values
        data={'分组':[group]*n,
              '性别':categories('性别',['男','女'],[g['male'],100-g['male']]),
              '年龄':_ranges(n,g['age'],0,rng),
              '病程':_ranges(n,g['course'],g['course_digits'],rng),
              '病情程度':categories('病情程度',['轻度','中度','重度'],g['severity']),
              '总有效率':categories('总有效率',['有效','无效'],[g['effective'],100-g['effective']])}
        if g.get('age_stats'):
            age_spec=dict(g['age_stats'],digits=0,parity='any',lower=g['age'][0]['lower'],upper=g['age'][0]['upper'])
            data['年龄']=_age_values(age_spec,n,rng,cancelled)
            actual_mean,actual_sd=statistics(data['年龄'])
            checks.append(dict(group=group,field='年龄',n=n,target_mean=age_spec['mean'],actual_mean=actual_mean,mean_error=actual_mean-age_spec['mean'],target_sd=age_spec['sd'],actual_sd=actual_sd,sd_error=actual_sd-age_spec['sd'],minimum=min(data['年龄']),maximum=max(data['年龄'])))
        for spec in (s for s in config.get('categoricals',[]) if s['group']==group):
            data[spec['name'].strip()]=categories(spec['name'],spec['labels'],spec['weights'])
        mode=g.get('extent_mode','category')
        if mode=='blank':data['病变范围']=['']*n
        elif mode=='range':data['病变范围']=_ranges(n,g['extent_ranges'],g.get('extent_digits',0),rng)
        else:data['病变范围']=categories('病变范围',[r['label'] for r in g['extent']],[r['percent'] for r in g['extent']])
        if g.get('recurrence_scope','all')=='effective':
            eligible=[i for i,v in enumerate(data['总有效率']) if v=='有效']
            values=_categorical(len(eligible),['复发','未复发'],[g['recurrence'],100-g['recurrence']],rng)
            data['复发率']=['未复发']*n
            for i,v in zip(eligible,values):data['复发率'][i]=v
            count=values.count('复发');denominator=len(eligible)
            checks.append(dict(group=group,field='复发率',category='复发（有效者为分母）',n=denominator,target_percent=g['recurrence'],actual_count=count,actual_percent=count/denominator*100 if denominator else None))
        else:data['复发率']=categories('复发率',['复发','未复发'],[g['recurrence'],100-g['recurrence']])
        for field,key in [('年龄','age'),('病程','course')]:
            for (r,lo,hi),count in zip(_effective_ranges(g[key],0 if key=='age' else g['course_digits']),allocation(n,[r['percent'] for r in g[key]])):
                factor=10**(0 if key=='age' else g['course_digits'])
                actual=sum(lo<=round(v*factor)<=hi for v in data[field])
                checks.append(dict(group=group,field=field,category=f"{r['lower']:g}～{r['upper']:g}"+('（上界归入下一段）' if hi<r['upper']*factor else ''),target_percent=r['percent'],actual_count=actual,actual_percent=actual/n*100))
        for spec in (s for s in config['indicators'] if s['group']==group):
            if cancelled():
                raise InterruptedError('已取消生成，未保存新的结果。')
            progress(f"{group} · {spec['name']}")
            try:
                values=_indicator(spec,n,rng,cancelled)
            except ValueError as exc:
                raise ValueError(f"{group} / {spec['name']}：{exc}") from exc
            if cancelled():
                raise InterruptedError('已取消生成，未保存新的结果。')
            data[spec['name'].strip()]=values
            mean,sd=statistics(values)
            checks.append(dict(group=group,field=spec['name'],n=n,target_mean=spec['mean'],actual_mean=mean,mean_error=mean-spec['mean'],target_sd=spec['sd'],actual_sd=sd,sd_error=sd-spec['sd'],minimum=min(values),maximum=max(values)))
        from .typed_field_engine import apply_fields
        def typed_numeric(spec,count,random_state,cancel):
            # Preserve the existing bounded integer-age refinement for prose-table fields.
            fn=_age_values if spec.get('name')=='年龄' and spec.get('digits')==0 and spec.get('lower') is not None and spec.get('upper') is not None else _indicator
            return fn(spec,count,random_state,cancel)
        overridden=apply_fields(config,group,n,data,checks,rng,cancelled,typed_numeric,categories,statistics)
        if '总有效率' in overridden and '复发率' not in overridden and g.get('recurrence_scope','all')=='effective':
            checks[:]=[check for check in checks if not (check['group']==group and check['field']=='复发率')]
            eligible=[i for i,value in enumerate(data['总有效率']) if value=='有效']
            values=_categorical(len(eligible),['复发','未复发'],[g['recurrence'],100-g['recurrence']],rng)
            data['复发率']=['未复发']*n
            for i,value in zip(eligible,values):data['复发率'][i]=value
            denominator=len(eligible)
            checks.append(dict(group=group,field='复发率',category='复发（有效者为分母）',n=denominator,target_percent=g['recurrence'],actual_count=values.count('复发'),actual_percent=values.count('复发')/denominator*100 if denominator else None))
        blank_fields=[field for field in g.get('blank_fields',[]) if field not in overridden]
        for field in blank_fields:data[field]=['']*n
        checks[:]=[check for check in checks if check['group']!=group or (check['field'] not in blank_fields and check['field'] in columns)]
        start=len(rows)+1
        rows.extend([[start+i]+[data[c][i] for c in columns[1:]] for i in range(n)])
    return dict(columns=columns,rows=rows,checks=checks,config=config)


def export_batch(result,path):
    """Write atomically: original files survive failed/cancelled saves."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    path=Path(path)
    book=Workbook();sheet=book.active;sheet.title='模拟病例数据'
    sheet.append(result['columns'])
    for row in result['rows']:
        sheet.append(row)
    groups={g['name']:g for g in result['config']['groups']}
    specs={(v['group'],v['name'].strip()):v for v in result['config']['indicators']}
    specs.update({(v['group'],v['name'].strip()):v for v in result['config'].get('typed_fields',[]) if v['kind'] in ('mean_sd','quantile')})
    for row in sheet.iter_rows(min_row=2):
        group=next(iter(groups.values())) if result['config'].get('ungrouped') else groups[row[result['columns'].index('分组')].value]
        for col,cell in enumerate(row):
            if not isinstance(cell.value,(int,float)):continue
            name=result['columns'][col]
            digits=specs[(group['name'],name)]['digits'] if (group['name'],name) in specs else group['course_digits'] if name=='病程' else group.get('extent_digits',0) if name=='病变范围' else 0
            cell.number_format='0'+('.'+'0'*digits if digits else '')
    audit=book.create_sheet('统计核对')
    fields=[('group','分组'),('field','字段'),('category','类别或区间'),('n','例数'),('target_mean','目标均数'),('actual_mean','实际均数'),('mean_error','均数偏差'),('target_sd','目标标准差'),('actual_sd','实际标准差'),('sd_error','标准差偏差'),('minimum','最小值'),('maximum','最大值'),('target_percent','目标比例%'),('actual_count','实际人数'),('actual_percent','实际比例%'),('target_q1','目标Q1'),('actual_q1','实际Q1'),('target_median','目标中位数'),('actual_median','实际中位数'),('target_q3','目标Q3'),('actual_q3','实际Q3')]
    audit.append(['数据集' if k=='group' and result['config'].get('ungrouped') else v for k,v in fields])
    for check in result['checks']:
        audit.append([check.get(k) for k,v in fields])
    settings=book.create_sheet('生成配置');settings.append(['设置','值'])
    settings.append(['数据性质','模拟数据，非原始病例；各指标独立生成，不还原配对或检验结果。'])
    settings.append(['比例处理','按配置比例换算人数后随机打乱；复发分母分别见各组 recurrence_scope（all=全组，effective=有效者）。'])
    settings.append(['病程单位',result['config'].get('course_unit','月')])
    for key,value in result['config'].items():
        # One indicator per row avoids Excel's 32,767 character cell limit.
        if key in ('indicators','groups','categoricals','field_rules','typed_fields'):
            for i,item in enumerate(value,1):
                settings.append([f'{key}[{i}]',json.dumps(item,ensure_ascii=False,allow_nan=False)])
        else:
            settings.append([key,json.dumps(value,ensure_ascii=False,allow_nan=False)])
    for ws in book:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='0D7668')
        for i in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(i)].width=20 if ws is not settings else (28 if i==1 else 100)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type='s'  # user-entered labels are text, never formulas
    fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=path.parent);os.close(fd)
    try:
        book.save(tmp);os.replace(tmp,path)
    finally:
        book.close()
        if os.path.exists(tmp):os.unlink(tmp)
    return path
