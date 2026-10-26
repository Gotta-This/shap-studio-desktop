"""Operate the restored Word workflow in source and bundled applications."""
from pathlib import Path
from collections import Counter
import json,time,traceback,sys
from PySide6.QtWidgets import QApplication,QFileDialog
from .word_batch_engine import BASE_COLUMNS,END_COLUMNS,export_batch


def verify_word(window,directory,source):
    root=Path(directory);root.mkdir(parents=True,exist_ok=True);app=QApplication.instance();checks=[]
    original_open=QFileDialog.getOpenFileName;original_save=QFileDialog.getSaveFileName
    def check(name,ok):
        checks.append(dict(name=name,passed=bool(ok)))
        if not ok:raise AssertionError(name)
    def pump():app.processEvents()
    def shot(name):pump();window.grab().save(str(root/(name+'.png')))
    def wait(tool):
        deadline=time.monotonic()+120
        while tool.is_running() and time.monotonic()<deadline:pump();time.sleep(.01)
        pump();check('后台任务结束',not tool.is_running())
    def rows(editor,data):
        editor.table.setRowCount(0)
        for row in data:editor.add(row)
    try:
        window.timer.stop();tool=window.open_module('word');pump()
        check('第三工具完整配置直接位于主界面',tool.isVisible() and tool.parentWidget() is window.boards[3])
        check('其他三个工具未实例化',set(window.modules)=={'word'})
        tool.common_n.setValue(70);first=tool.group_editors[0]
        first.fields['male'].setText('60');first.fields['light'].setText('20');first.fields['medium'].setText('50');first.fields['severe'].setText('30')
        rows(first.age,[(10,20,40),(20,30,60)]);rows(first.course,[(1,3,50),(3,6,50)]);first.course_digits.setValue(1)
        first.extent_mode.setCurrentIndex(first.extent_mode.findData('category'));rows(first.extent,[('局限',40),('广泛',60)])
        first.fields['effective'].setText('80');first.fields['recurrence'].setText('10')
        second=tool.group_editors[1];second.fields['effective'].setText('90');second.fields['recurrence'].setText('20');second.recurrence_scope.setCurrentIndex(second.recurrence_scope.findData('effective'))
        shot('01-basic-workflow')
        QFileDialog.getOpenFileName=lambda *a,**kw:(str(source),'Word');tool.import_button.click();pump();tool.group_inputs[0].setText('对照组');tool.add_group_input('中药组');tool.confirm_groups();pump()
        check('实际Word导入30个前后指标',tool.metrics.rowCount()==60)
        tool.leave_unknown_blank.setChecked(False)  # Explicitly use the manually entered basic-data rules.
        check('导入保留用户70例和80%有效设置',tool.group_editors[0].fields['n'].text()=='70' and float(tool.group_editors[0].fields['effective'].text())==80)
        config=tool.get_config();check('两组共用性别年龄病程配置',config['groups'][0]['age']==config['groups'][1]['age'] and config['groups'][1]['male']==60)
        check('两组有效和复发比例独立',config['groups'][1]['effective']==90 and config['groups'][1]['recurrence']==20)
        shot('02-word-all-indicators')
        path=root/'configuration.json';QFileDialog.getSaveFileName=lambda *a,**kw:(str(path),'JSON');tool.save_button.click()
        check('完整配置保存',json.loads(path.read_text(encoding="utf-8"))==config)
        tool.seed.setText('99');QFileDialog.getOpenFileName=lambda *a,**kw:(str(path),'JSON');tool.load_button.click();check('完整配置载入往返一致',tool.get_config()==config)
        tool.generate_button.click();wait(tool);check('实际按钮生成成功',tool.result is not None)
        r=tool.result;check('140例和39列',len(r['rows'])==140 and len(r['columns'])==39)
        check('固定前7列及后2列',r['columns'][:7]==BASE_COLUMNS and r['columns'][-2:]==END_COLUMNS)
        check('编号连续1至140',[row[0] for row in r['rows']]==list(range(1,141)))
        check('两组各70例',Counter(row[1] for row in r['rows'])=={'对照组':70,'中药组':70})
        check('每组性别比例42男28女',all(Counter(row[2] for row in r['rows'] if row[1]==g)=={'男':42,'女':28} for g in ('对照组','中药组')))
        check('年龄相邻区间按28和42人分配',all(sum(10<=row[3]<20 for row in r['rows'] if row[1]==g)==28 and sum(20<=row[3]<=30 for row in r['rows'] if row[1]==g)==42 for g in ('对照组','中药组')))
        check('病变范围按28和42人分配',all(Counter(row[6] for row in r['rows'] if row[1]==g)=={'局限':28,'广泛':42} for g in ('对照组','中药组')))
        check('结果值严格为有效无效及复发未复发',set(row[-2] for row in r['rows'])=={'有效','无效'} and set(row[-1] for row in r['rows'])=={'复发','未复发'})
        second_rows=[row for row in r['rows'] if row[1]=='中药组'];check('中药组无效者不参与有效者分母的复发生成',all(row[-1]=='未复发' for row in second_rows if row[-2]=='无效'))
        check('60个组别指标的均数标准差偏差可核对',len([v for v in r['checks'] if 'actual_mean' in v])==60 and all(abs(v['mean_error'])<.02 and abs(v['sd_error'])<.02 for v in r['checks'] if 'actual_mean' in v))
        tool.copy_button.click();check('复制包含完整表头和140行',len(app.clipboard().text().splitlines())==141)
        excel=root/'模拟病例数据.xlsx';QFileDialog.getSaveFileName=lambda *a,**kw:(str(excel),'Excel');tool.export_button.click();wait(tool)
        from openpyxl import load_workbook
        book=load_workbook(excel,data_only=True);check('导出三个工作表',book.sheetnames==['模拟病例数据','统计核对','生成配置'])
        data=list(book.worksheets[0].values);check('Excel每个单元格与生成结果一致',list(data[0])==r['columns'] and [list(row) for row in data[1:]]==r['rows']);book.close()
        shot('03-generated-cases-and-export')
        before=r;window.open_module('category');window.open_module('word');check('切换工具保留第三模块参数和结果',tool.result is before and tool.get_config()==config)
        report=dict(passed=True,checks=checks,source=str(source),runtime='bundled' if getattr(sys,'frozen',False) else 'source',rows=140,columns=r['columns'],numeric_checks=[v for v in r['checks'] if 'actual_mean' in v]);(root/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8");return True
    except Exception:
        (root/'verification.json').write_text(json.dumps(dict(passed=False,checks=checks,traceback=traceback.format_exc()),ensure_ascii=False,indent=2),encoding="utf-8");return False
    finally:QFileDialog.getOpenFileName=original_open;QFileDialog.getSaveFileName=original_save
