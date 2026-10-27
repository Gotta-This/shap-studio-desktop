"""Neutral binary defaults preserve reported counts; invalid generation shows a modal."""
import os,tempfile,unittest
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from pathlib import Path
from PySide6.QtCore import Qt,QCoreApplication,QEvent
from clinical_studio.app import create_application
from clinical_studio.strict_ui import DesktopWindow
from clinical_studio.editor_dialogs import EditorDialog
from clinical_studio.word_batch_engine import read_word,default_batch_config,generate_batch,export_batch
from clinical_studio.word_table_types import summary
from clinical_studio.typed_field_ui import convert
from test_word_table_types import save
APP=create_application([])
class BinaryDefaultTests(unittest.TestCase):
 def test_consistent_counts_use_numeric_labels_and_export_numbers(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'table.docx';save(p,[[['组别','例数','止血成功','住院期间再出血'],['对照组','15','12（80.00%）','4（26.67%）'],['试验组','15','15（100.00%）','1（6.67%）']]])
   fields=read_word(p,['对照组','试验组']).typed_fields
   self.assertTrue(all(f['kind']=='binary01' and f['labels']==[1,0] and f['confirmed'] for f in fields))
   c=default_batch_config(['对照组','试验组']);c['include_basics']=False;c['typed_fields']=fields
   for g in c['groups']:g['n']=15
   r=generate_batch(c);idx=r['columns'].index('止血成功');v=[row[idx] for row in r['rows'] if row[1]=='对照组'];self.assertEqual(v.count(1),12);self.assertEqual(v.count(0),3)
   out=Path(tmp)/'cases.xlsx';export_batch(r,out)
   from openpyxl import load_workbook
   book=load_workbook(out);self.assertIsInstance(book.worksheets[0].cell(2,idx+1).value,int);book.close()
 def test_inconsistent_missing_n_and_aggregate_remain_pending(self):
  for raw,n in [('4(40.00%)',15),('4(26.67%)',None),('16(106.67%)',15)]:
   f=summary('止血成功','对照组',raw,n,1);self.assertFalse(f['confirmed']);self.assertFalse(convert(f,f['kind'])['confirmed'])
  self.assertFalse(summary('总发生率','对照组','4(26.67%)',15,1)['confirmed'])
 def test_clear_complication_and_gender_labels_are_kept(self):
  self.assertEqual(summary('恶心呕吐','对照组','4(16.00%)',25,1,table_title='表4 并发症')['labels'],['有','无'])
  self.assertEqual(summary('性别（男/女）','对照组','9/6',15,1)['labels'],['男','女'])
 def test_invalid_generate_opens_one_native_dialog_and_preserves_bottom_message(self):
  w=DesktopWindow();w.timer.stop();w.show();t=w.open_module('word');APP.processEvents()
  try:
   t.add_indicator();t.metrics.item(0,3).setText('');t.generate();APP.processEvents();d=t.validation_dialog
   self.assertIsInstance(d,EditorDialog);self.assertIs(d.parentWidget(),w);self.assertIsNone(d.graphicsProxyWidget());self.assertTrue(d.isVisible());self.assertEqual(d.windowModality(),Qt.WindowModal)
   self.assertTrue(t.message.text());self.assertFalse(t.is_running());d.accept();APP.processEvents()
   t.show_generation_error('第二次错误');APP.processEvents();self.assertIsNot(t.validation_dialog,d);self.assertEqual(t.message.text(),'第二次错误');t.validation_dialog.accept()
  finally:w.close();w.deleteLater();QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)
