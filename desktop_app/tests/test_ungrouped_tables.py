"""Unnamed vertical datasets are parsed, editable and exported without fake groups."""
import os,tempfile,unittest,json
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from pathlib import Path
from unittest.mock import patch
from PySide6.QtWidgets import QFileDialog
from clinical_studio.app import create_application
from clinical_studio.word_batch import WordBatchTool
from clinical_studio.ungrouped_tables import parse_ungrouped,DATASET
from clinical_studio.merged_table_verify import save_fixture
from clinical_studio.word_batch_engine import generate_batch,export_batch
APP=create_application([])
ROWS=[['项目','数值（n=20）'],['性别（男/女）','18(90)/2(10)'],['年龄（岁）','58.35±10.41'],['BMI（kg/m²）','23.85±2.74'],['合并高血压','7(35)']]

class UngroupedTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=save_fixture(Path(self.tmp.name)/'one.docx',[ROWS]);self.tool=WordBatchTool();self.tool.show();APP.processEvents()
    def tearDown(self):self.tool.close();self.tool.deleteLater();APP.processEvents();self.tmp.cleanup()
    def test_no_group_names_are_required_for_preview_and_apply(self):
        t=self.tool;t.import_word(self.path)
        self.assertEqual(t.word_scan['candidates'],[]);self.assertTrue(t.ungrouped_option.isChecked());self.assertEqual(t.ungrouped_n.value(),20)
        self.assertEqual(len(t.preview_draft['fields']),4);self.assertEqual(t.group_inputs[0].text(),'');self.assertFalse(t.generate_button.isEnabled())
        t.confirm_groups();self.assertTrue(t.groups_confirmed,t.message.text());self.assertTrue(t.ungrouped);self.assertEqual(len(t.typed_specs),4)
    def test_generated_excel_contains_only_id_and_source_fields(self):
        t=self.tool;t.import_word(self.path);t.confirm_groups();config=t.get_config();self.assertTrue(config['ungrouped'])
        result=generate_batch(config);self.assertEqual(result['columns'],['编号','性别','年龄（岁）','BMI（kg/m²）','合并高血压'])
        self.assertEqual(len(result['rows']),20);self.assertEqual(sum(r[1]=='男' for r in result['rows']),18)
        self.assertEqual(sum(r[-1]=='有' for r in result['rows']),7)
        out=Path(self.tmp.name)/'cases.xlsx';export_batch(result,out)
        from openpyxl import load_workbook
        book=load_workbook(out);self.assertEqual(book.worksheets[0].max_row,21);self.assertNotIn('分组',[c.value for c in book.worksheets[0][1]])
        self.assertEqual(book.worksheets[0].cell(2,3).number_format,'0.00');self.assertEqual(book['统计核对'].cell(1,1).value,'数据集');book.close()
    def test_missing_sample_size_allows_recognition_but_requires_input_for_apply(self):
        save_fixture(self.path,[[['指标','数值'],['年龄','60±5'],['BMI','23±2']]])
        t=self.tool;t.import_word(self.path);self.assertTrue(t.ungrouped_option.isChecked());self.assertEqual(len(t.preview_draft['fields']),2);self.assertEqual(t.ungrouped_n.value(),0)
        t.confirm_groups();self.assertFalse(t.groups_confirmed);self.assertIn('例数',t.message.text())
        t.ungrouped_n.setValue(30);t.confirm_groups();self.assertTrue(t.groups_confirmed);self.assertEqual(len(generate_batch(t.get_config())['rows']),30)
    def test_generation_size_change_preserves_original_source_proportions(self):
        t=self.tool;t.import_word(self.path);t.ungrouped_n.setValue(30);t.confirm_groups()
        result=generate_batch(t.get_config());self.assertEqual(len(result['rows']),30)
        self.assertEqual(sum(r[1]=='男' for r in result['rows']),27)
        self.assertEqual(next(f for f in t.typed_specs if f['name']=='性别')['source_n'],20)
    def test_preview_parameter_and_classification_edits_survive_apply(self):
        t=self.tool;t.import_word(self.path);f=next(f for f in t.preview_draft['fields'] if f['name']=='年龄（岁）');f['mean']=61.23;f['classification']='基线'
        t.confirm_groups();f=next(f for f in t.typed_specs if f['name']=='年龄（岁）');self.assertEqual(f['mean'],61.23);self.assertEqual(f['classification'],'基线')
    def test_category_subrows_and_three_column_category_table(self):
        for rows in ([['项目','数值'],['样本量','20'],['性别',''],['男','18(90)'],['女','2(10)'],['年龄','60±5']],
                     [['项目','类别','n(%)'],['性别','男','18(90)'],['','女','2(10)']]):
            parsed=parse_ungrouped(rows,1);f=next(f for f in parsed['fields'] if f['name']=='性别')
            self.assertEqual(f['labels'],['男','女']);self.assertEqual(f['weights'],[90,10]);self.assertTrue(f['confirmed']);self.assertEqual(parsed['n'],20)
    def test_complication_subrows_are_independent_binary_fields(self):
        rows=[['项目','总体(n=20)'],['并发症',''],['尿潴留','2'],['切口感染','1']]
        parsed=parse_ungrouped(rows,1)
        self.assertEqual({f['name'] for f in parsed['fields']},{'尿潴留','切口感染'})
        self.assertTrue(all(f['kind']=='yesno' and f['labels']==['有','无'] for f in parsed['fields']))
        self.assertEqual(parsed['fields'][0]['weights'],[10,90])
        rows[0]=['项目','数值'];parsed=parse_ungrouped(rows,1)
        self.assertIsNone(parsed['n']);self.assertTrue(all(not f['confirmed'] for f in parsed['fields']))
    def test_custom_named_groups_are_not_mistaken_for_ungrouped_data(self):
        parsed=parse_ungrouped([['项目','甲队','乙队'],['年龄','60±5','61±4']],1)
        self.assertFalse(parsed['handled'])
        parsed=parse_ungrouped([['组别','评分'],['甲队','60±5'],['乙队','61±4']],1);self.assertFalse(parsed['handled'])
    def test_conflicting_sample_sizes_require_manual_review(self):
        save_fixture(self.path,[ROWS,[['项目','总体(n=30)'],['其他指标','10±2']]])
        t=self.tool;t.import_word(self.path);self.assertEqual(t.ungrouped_n.value(),0);self.assertIn('例数不同',t.match_report.toPlainText());self.assertFalse(t.confirm_groups_button.isEnabled())
    def test_save_reload_preserves_ungrouped_mode(self):
        t=self.tool;t.import_word(self.path);t.confirm_groups();config=t.get_config();path=Path(self.tmp.name)/'config.json';path.write_text(json.dumps(config,ensure_ascii=False))
        with patch.object(QFileDialog,'getOpenFileName',return_value=(str(path),'JSON')):t.load_config()
        self.assertEqual(t.get_config(),config);self.assertTrue(t.ungrouped);self.assertNotIn('分组',generate_batch(t.get_config())['columns'])
    def test_reimport_grouped_document_resets_ungrouped_mode(self):
        t=self.tool;t.import_word(self.path);t.confirm_groups()
        save_fixture(self.path,[[['组别','例数','评分'],['对照组','20','10±2'],['试验组','20','11±2']]])
        t.import_word(self.path);self.assertFalse(t.ungrouped);self.assertFalse(t.ungrouped_option.isChecked());t.use_candidate_groups();t.confirm_groups();self.assertIn('分组',generate_batch(t.get_config())['columns'])

if __name__=='__main__':unittest.main()
