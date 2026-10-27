"""Selected Word group matching, arbitrary group generation and editable columns."""
import os,tempfile,unittest
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from pathlib import Path
from collections import Counter
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QLineEdit
from PySide6.QtTest import QTest
from clinical_studio.app import create_application
from clinical_studio.strict_ui import DesktopWindow
from clinical_studio.word_batch_engine import W,read_word,default_batch_config,generate_batch,validate_batch,export_batch
APP=create_application([])

def document(path):
    root=ET.Element(W+'document');body=ET.SubElement(root,W+'body')
    tables=[
      [('编号','组别','例数','治疗前','治疗后'),('1','常规 治疗组','8','20±2','15±2'),('2','观察组','9','21±2','14±2'),('3','观察组1','10','22±2','13±2')],
      [('组别','氧合指数'),('常规治疗组','202.64±10'),('观察组','210.42±11'),('观察组1','215.42±12')],
      [('说明','无关表格'),('备注','不含组别')]]
    for rows in tables:
        table=ET.SubElement(body,W+'tbl')
        for values in rows:
            row=ET.SubElement(table,W+'tr')
            for value in values:
                cell=ET.SubElement(row,W+'tc');p=ET.SubElement(cell,W+'p');r=ET.SubElement(p,W+'r');ET.SubElement(r,W+'t').text=value
    with ZipFile(path,'w') as z:z.writestr('word/document.xml',ET.tostring(root))

class WordGroupTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/'custom.docx';document(self.path)
        self.w=DesktopWindow();self.w.timer.stop();self.w.show();self.t=self.w.open_module('word');APP.processEvents()
    def tearDown(self):self.w.close();self.w.deleteLater();APP.processEvents();self.tmp.cleanup()
    def select(self,names):
        self.t.import_word(self.path);self.t.show_group_selection(names)
        self.t.confirm_groups();APP.processEvents()
    def test_import_waits_for_explicit_confirmation_and_counts_all_tables(self):
        t=self.t;t.import_word(self.path)
        self.assertFalse(t.groups_confirmed);self.assertTrue(t.group_panel.isVisible());self.assertEqual([e.text() for e in t.group_inputs],['常规治疗组','观察组','观察组1'])
        self.assertFalse(t.generate_button.isEnabled());self.assertEqual(t.metrics.rowCount(),0)
        t.show_group_selection(['观察组','常规治疗组']);r=t.match_groups()
        self.assertEqual(r.table_count,3);self.assertEqual(r.group_matches,{'观察组':2,'常规治疗组':2})
        self.assertEqual(r.indicators,[]);self.assertIn('2/2',t.match_summary.text());self.assertIn('表3：未匹配',t.match_report.toPlainText())
        t.confirm_groups();self.assertTrue(t.groups_confirmed);self.assertEqual(t.metrics.rowCount(),6)
        self.assertEqual([e.name for e in t.group_editors],['观察组','常规治疗组']);self.assertIn('合计140例',t.total_label.text());self.assertIn('观察组',t.group_selector.itemText(0))
    def test_full_name_matching_missing_and_duplicate_groups(self):
        r=read_word(self.path,['观察组']);self.assertEqual(r.group_matches,{'观察组':2});self.assertEqual(len(r.indicators),3)
        t=self.t;t.import_word(self.path);t.show_group_selection(['不存在组']);t.confirm_groups();self.assertFalse(t.groups_confirmed);self.assertIn('未在Word',t.message.text())
        t.show_group_selection(['观察组','观 察组']);self.assertIsNone(t.match_groups());self.assertIn('重复',t.match_summary.text())
    def test_single_and_three_groups_generate_and_export(self):
        for names in [['观察组'],['观察组','常规治疗组','观察组1']]:
            r=read_word(self.path,names);c=default_batch_config(names);c['indicators']=r.indicators
            for g in c['groups']:g['n']=8
            result=generate_batch(c);self.assertEqual(Counter(row[1] for row in result['rows']),{n:8 for n in names})
            out=Path(self.tmp.name)/'out.xlsx';export_batch(result,out)
            from openpyxl import load_workbook
            b=load_workbook(out);self.assertEqual(b.worksheets[0].max_row,8*len(names)+1);b.close()
            if len(names)>1:
                c['indicators'].pop()
                with self.assertRaises(ValueError):validate_batch(c)
    def test_actual_column_editor_syncs_all_groups_and_export(self):
        self.select(['观察组','常规治疗组','观察组1']);t=self.t
        field=t.metrics.item(0,1);self.assertTrue(field.flags()&Qt.ItemIsEditable)
        t.metrics.editItem(field);APP.processEvents();editor=t.metrics.findChild(QLineEdit)
        self.assertIsNotNone(editor);editor.setText('治疗前自定义名称');QTest.keyClick(editor,Qt.Key_Return);APP.processEvents()
        self.assertEqual([t.metrics.item(i,1).text() for i in range(3)],['治疗前自定义名称']*3)
        t.common_n.setValue(8);c=t.get_config();result=generate_batch(c);self.assertIn('治疗前自定义名称',result['columns'])
        t.shared_basics.setChecked(False);self.assertEqual(len(t.get_config()['groups']),3)
    def test_three_group_shared_basics_and_config_roundtrip(self):
        self.select(['观察组','常规治疗组','观察组1']);t=self.t;t.add_indicator()
        self.assertEqual(t.metrics.rowCount(),12);t.group_editors[0].fields['male'].setText('67');t.group_editors[2].fields['effective'].setText('91')
        c=t.get_config();self.assertEqual([g['male'] for g in c['groups']],[67]*3);self.assertEqual(c['groups'][2]['effective'],91)
        from unittest.mock import patch
        from PySide6.QtWidgets import QFileDialog
        p=Path(self.tmp.name)/'config.json'
        with patch.object(QFileDialog,'getSaveFileName',return_value=(str(p),'JSON')):t.save_config()
        t.edit_groups();self.assertFalse(t.groups_confirmed)
        with patch.object(QFileDialog,'getOpenFileName',return_value=(str(p),'JSON')):t.load_config()
        self.assertEqual(t.get_config(),c);self.assertTrue(t.groups_confirmed)

    def test_new_word_clears_previous_results_metrics_and_group_rules(self):
        from unittest.mock import patch
        from PySide6.QtWidgets import QFileDialog
        self.select(['观察组','常规治疗组']);t=self.t;t.common_n.setValue(8)
        groups=t.get_config()['groups']
        for group in groups:
            group.update(age=[dict(lower=99,upper=99,percent=100.)],effective=99.,
                         extent_mode='category',extent=[dict(label='旧文档专属类别',percent=100.)])
        t.set_groups(groups);t.add_indicator()
        t.show_result(generate_batch(t.get_config()))
        self.assertIsNotNone(t.result);self.assertTrue(t.export_button.isEnabled())
        other=Path(self.tmp.name)/'new.docx'
        with ZipFile(self.path) as old,ZipFile(other,'w') as new:
            root=ET.fromstring(old.read('word/document.xml'))
            for text in root.iter(W+'t'):
                if text.text=='氧合指数':text.text='新文档指标'
            new.writestr('word/document.xml',ET.tostring(root))
        t.import_word(other)
        self.assertIsNone(t.result);self.assertEqual(t.metrics.rowCount(),0)
        self.assertEqual(t.preview.rowCount(),0);self.assertEqual(t.preview.columnCount(),0)
        self.assertFalse(t.copy_button.isEnabled());self.assertFalse(t.export_button.isEnabled())
        with patch.object(QFileDialog,'getSaveFileName') as save:
            t.export_data();save.assert_not_called()
        t.show_group_selection(['观察组','常规治疗组']);t.confirm_groups()
        config=t.get_config();names={s['name'] for s in config['indicators']}
        self.assertIn('新文档指标',names);self.assertNotIn('氧合指数',names)
        self.assertFalse(any(name.startswith('自定义指标') for name in names))
        for group in config['groups']:
            self.assertEqual(group['extent_mode'],'blank')
            self.assertNotEqual(group['age'],groups[0]['age']);self.assertEqual(group['effective'],80.)
        result=generate_batch(config)
        self.assertTrue(all(row[6]=='' for row in result['rows']))
        self.assertEqual(result['config']['source'],str(other.resolve()))

    def test_editing_groups_in_same_word_keeps_rules_but_invalidates_result(self):
        self.select(['观察组','常规治疗组']);t=self.t;t.common_n.setValue(8)
        t.group_editors[0].fields['effective'].setText('63')
        t.show_result(generate_batch(t.get_config()))
        t.edit_groups();self.assertIsNone(t.result);t.confirm_groups()
        self.assertEqual(t.get_config()['groups'][0]['effective'],63.)
if __name__=='__main__':unittest.main()
