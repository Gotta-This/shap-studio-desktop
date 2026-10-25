"""Word table import, per-group synthesis and native GUI/export regression."""
import os
import tempfile
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
os.environ.setdefault('MPLCONFIGDIR',os.path.join(tempfile.gettempdir(),'dbq-matplotlib-check'))
from collections import Counter
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import time
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from clinical_studio.word_batch_engine import (
    W, BASE_COLUMNS, END_COLUMNS, GROUPS, read_word, default_batch_config,
    allocation, validate_batch, generate_batch, export_batch,
)


def fixture_docx(path, missing=False):
    root=ET.Element(W+'document');body=ET.SubElement(root,W+'body');tbl=ET.SubElement(body,W+'tbl')
    def row(cells):
        tr=ET.SubElement(tbl,W+'tr')
        for text,span in cells:
            tc=ET.SubElement(tr,W+'tc')
            if span>1:
                props=ET.SubElement(tc,W+'tcPr');ET.SubElement(props,W+'gridSpan',{W+'val':str(span)})
            p=ET.SubElement(tc,W+'p');r=ET.SubElement(p,W+'r');ET.SubElement(r,W+'t').text=text
    row([('组别',1),('例数',1),('主要症状',3)])
    row([('',1),('',1),('治疗前',1),('治疗后',2)])
    row([('对照组',1),('10',1),('16.98±2.43',1),('9.05±1.65*',2)])
    row([('中药组',1),('12',1),('' if missing else '16.36±2.57',1),('7.51±1.43*',2)])
    row([('P',1),('',1),('0.155',1),('＜0.001',2)])
    with ZipFile(path,'w') as z:z.writestr('word/document.xml',ET.tostring(root))


def small_config():
    c=default_batch_config()
    for i,g in enumerate(c['groups']):
        g.update(n=10+i*2,male=50.,severity=[30.,50.,20.],age=[dict(lower=18,upper=29,percent=40.),dict(lower=30,upper=60,percent=60.)],course=[dict(lower=1,upper=12,percent=100.)])
        c['indicators'].append(dict(table=1,name='前-测试',group=g['name'],mean=10,sd=2,digits=2,lower=0,upper=20,parity='any',raw='10±2'))
    return c


class WordBatchTests(unittest.TestCase):
    def test_merged_grid_and_missing_values(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'table.docx';fixture_docx(p);r=read_word(p)
            self.assertEqual(len(r.indicators),4)
            self.assertEqual([s['name'] for s in r.indicators[::2]],['前-主要症状','后-主要症状'])
            self.assertEqual(r.group_sizes,{'对照组':10,'中药组':12})
            self.assertEqual(r.indicators[0]['mean'],16.98)
            self.assertEqual(r.indicators[2]['sd'],1.65)
            fixture_docx(p,missing=True);r=read_word(p)
            self.assertIsNone(r.indicators[1]['mean'])
            self.assertTrue(any('未识别均数' in m for m in r.messages))

    def test_allocation_validation_and_cancellation(self):
        self.assertEqual(allocation(70,[50,50]),[35,35])
        self.assertEqual(sum(allocation(7,[33,34,33])),7)
        for weights in [[40,40],[-1,101],[float('nan'),100]]:
            with self.assertRaises(ValueError):allocation(70,weights)
        c=small_config();validate_batch(c)
        for mutate in [lambda c:c['groups'][0].update(severity=[40,40,0]),lambda c:c['groups'][0]['age'].append(dict(lower=20,upper=40,percent=0)),lambda c:c['indicators'].pop(),lambda c:c['indicators'][0].update(sd=-1),lambda c:c.update(seed=-1)]:
            broken=deepcopy(c);mutate(broken)
            with self.assertRaises(ValueError):validate_batch(broken)
        with self.assertRaises(InterruptedError):generate_batch(c,cancelled=lambda:True)

    def test_generate_repeatability_and_excel(self):
        c=small_config();r=generate_batch(c);again=generate_batch(c)
        self.assertEqual(r,again)
        self.assertEqual(r['columns'],BASE_COLUMNS+['前-测试']+END_COLUMNS)
        self.assertEqual(len(r['rows']),22)
        self.assertEqual([row[0] for row in r['rows']],list(range(1,23)))
        self.assertEqual(Counter(row[1] for row in r['rows']),{'对照组':10,'中药组':12})
        self.assertEqual(Counter(row[2] for row in r['rows'][:10]),{'男':5,'女':5})
        self.assertEqual(sum(18<=row[3]<=29 for row in r['rows'][:10]),4)
        for check in r['checks']:
            if 'actual_mean' in check:
                self.assertLess(abs(check['mean_error']),.1);self.assertLess(abs(check['sd_error']),.1)
        for parity in ['odd','even']:
            spec=dict(c['indicators'][0],parity=parity,digits=0,mean=11 if parity=='odd' else 10)
            cfg=deepcopy(c);cfg['indicators']=[dict(spec,group=g) for g in GROUPS]
            values=[row[7] for row in generate_batch(cfg)['rows']]
            self.assertTrue(all(x%2==(1 if parity=='odd' else 0) for x in values))
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'result.xlsx';export_batch(r,path)
            from openpyxl import load_workbook
            w=load_workbook(path,data_only=False)
            self.assertEqual(w.sheetnames,['模拟病例数据','统计核对','生成配置'])
            self.assertEqual(w.worksheets[0].max_row,23)
            self.assertEqual([cell.value for cell in w.worksheets[0][1]],r['columns'])
            self.assertEqual(list(w.worksheets[0].values)[1],tuple(r['rows'][0]))
            self.assertEqual(w.worksheets[0].freeze_panes,'A2');w.close()
            # Formula-looking user labels are exported as plain text.
            r['rows'][0][6]='=1+1';export_batch(r,path);w=load_workbook(path)
            self.assertEqual(w.worksheets[0]['G2'].data_type,'s');w.close()

    def test_native_ui_flow(self):
        from clinical_studio.app import create_application
        from clinical_studio.hub import DesktopWindow
        from PySide6.QtWidgets import QFileDialog
        from unittest.mock import patch
        app=create_application([]);window=DesktopWindow();window.show();app.processEvents()
        tool=window.open_module('word')
        with tempfile.TemporaryDirectory() as d:
            try:
                path=Path(d)/'source.docx';fixture_docx(path)
                tool.import_word(path);self.assertEqual(tool.metrics.rowCount(),4)
                self.assertEqual(tool.group_editors[1].fields['n'].text(),'12')
                tool.generate_button.click();deadline=time.monotonic()+20
                while tool.is_running() and time.monotonic()<deadline:app.processEvents();time.sleep(.01)
                app.processEvents();self.assertFalse(tool.is_running());self.assertIsNotNone(tool.result,tool.message.text())
                self.assertEqual(len(tool.result['rows']),22);tool.copy_button.click()
                self.assertEqual(len(app.clipboard().text().splitlines()),23)
                out=Path(d)/'gui.xlsx'
                with patch.object(QFileDialog,'getSaveFileName',return_value=(str(out),'Excel')):tool.export_button.click()
                deadline=time.monotonic()+10
                while tool.is_running() and time.monotonic()<deadline:app.processEvents();time.sleep(.01)
                app.processEvents()
                self.assertTrue(out.exists(),tool.message.text())
                conf=Path(d)/'config.json'
                with patch.object(QFileDialog,'getSaveFileName',return_value=(str(conf),'JSON')):tool.save_button.click()
                self.assertTrue(conf.exists());self.assertEqual(json.loads(conf.read_text(encoding='utf-8'))['groups'][1]['n'],12)
                tool.seed.setText('99')
                with patch.object(QFileDialog,'getOpenFileName',return_value=(str(conf),'JSON')):tool.load_button.click()
                self.assertEqual(tool.seed.text(),'12')
                previous=tool.result;window.show_home();self.assertIs(window.open_module('word'),tool);self.assertIs(tool.result,previous)
            finally:
                if tool.worker:tool.worker.wait(20000)
                window.close();window.deleteLater();app.processEvents()

if __name__=='__main__':unittest.main()
